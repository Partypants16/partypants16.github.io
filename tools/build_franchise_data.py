"""Turn the franchise tracker spreadsheet (Book.xlsx) into franchise-data.js,
the data file that franchise.html shows.

Run it from the my-website folder after changing the spreadsheet:

    python tools/build_franchise_data.py              # reads Book.xlsx
    python tools/build_franchise_data.py other.xlsx

Needs openpyxl (pip install openpyxl).

How a sheet is read (each worksheet is one franchise):
- Row 2 is the header. Column A holds the titles. After that, every person has
  a status column (their name, e.g. "Rachel <3") and a thoughts column (a
  header containing "thoughts", e.g. "Rachel Thoughts"). The header ends at the
  first empty cell.
- A person's watch status is the fill colour of their status cell, looked up in
  the colour legend to the right of the header (a label cell with a filled cell
  next to it). A cell with no fill counts as "not marked yet".
- A bold title in column A, with no colours or thoughts in its row, starts a
  new section (e.g. "Original X-Men Movies"). Titles above the first one go in
  the sheet's first section, which is named in SHEETS below. Empty rows are
  skipped.
"""

import json
import re
import sys
from pathlib import Path

import openpyxl

HEADER_ROW = 2

# Settings per worksheet. A sheet that isn't listed is still read: its key is
# the sheet name in lowercase and its first section is called "Movies".
SHEETS = {
    "Sheet1": {"key": "marvel", "name": "Marvel", "first_section": "MCU Movies"},
}

# Legend labels (lowercase, letters only) -> the status keys franchise.js knows.
STATUS_KEYS = {
    "rewatched recently": "rewatched",
    "watched": "watched",
    "never watched": "never",
    "wont watch": "wont",
    "not out yet": "notout",
}

# A white fill is the same as no fill.
NO_FILL = {"rgb:FFFFFFFF", "rgb:00FFFFFF"}


def fill_key(cell):
    """A comparable key for a cell's fill colour, or None if it has no fill."""
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return None
    color = fill.fgColor
    if color.type == "rgb":
        key = "rgb:" + str(color.rgb).upper()
    elif color.type == "theme":
        key = "theme:%s:%s" % (color.theme, round(color.tint or 0, 3))
    else:
        key = "indexed:%s" % color.indexed
    return None if key in NO_FILL else key


def text(value):
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").strip()


def label_key(label):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z ]", "", label.lower())).strip()


def person_name(header):
    """'Rachel <3' / '<3 Paolo' / 'catherine thoughts' -> 'Rachel' / 'Paolo' / 'catherine'."""
    name = re.sub(r"thoughts?", "", header.replace("<3", ""), flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", name).strip()


def read_header(ws):
    """Returns (people, last header column); each person is a dict with
    name, status_col and thoughts_col (either column may be None)."""
    people = {}
    col = 2
    while text(ws.cell(HEADER_ROW, col).value):
        header = text(ws.cell(HEADER_ROW, col).value)
        name = person_name(header)
        person = people.setdefault(name.lower(), {"name": name, "status_col": None, "thoughts_col": None})
        if re.search(r"thoughts?", header, re.IGNORECASE):
            person["thoughts_col"] = col
        else:
            person["status_col"] = col
            person["name"] = name  # status headers have the nicer capitalisation
        col += 1
    return list(people.values()), col - 1


def read_legend(ws, first_col):
    """Fill colour key -> status key, from label cells with a filled cell to their right."""
    legend = {}
    for row in ws.iter_rows(min_col=first_col, max_col=ws.max_column):
        for cell in row[:-1]:
            label = text(cell.value)
            neighbour = ws.cell(cell.row, cell.column + 1)
            color = fill_key(neighbour)
            if not label or color is None or text(neighbour.value):
                continue
            key = STATUS_KEYS.get(label_key(label))
            if key is None:
                sys.exit(
                    'Unknown legend label "%s" at %s in sheet "%s". Add it to STATUS_KEYS in this '
                    "script and to STATUSES in franchise.js." % (label, cell.coordinate, ws.title)
                )
            legend[color] = key
    if not legend:
        sys.exit('No colour legend found to the right of the header in sheet "%s".' % ws.title)
    return legend


def read_sheet(ws):
    settings = SHEETS.get(ws.title, {})
    people, last_col = read_header(ws)
    legend = read_legend(ws, last_col + 2)

    sections = [{"name": settings.get("first_section", "Movies"), "titles": []}]
    for row in range(HEADER_ROW + 1, ws.max_row + 1):
        first = ws.cell(row, 1)
        title = text(first.value)
        if not title:
            continue

        status, thoughts = {}, {}
        for person in people:
            if person["status_col"]:
                cell = ws.cell(row, person["status_col"])
                color = fill_key(cell)
                if color is not None:
                    if color not in legend:
                        sys.exit(
                            "%s (%s, %s) has a colour that isn't in the legend."
                            % (cell.coordinate, title, person["name"])
                        )
                    status[person["name"]] = legend[color]
            if person["thoughts_col"]:
                thought = text(ws.cell(row, person["thoughts_col"]).value)
                if thought:
                    thoughts[person["name"]] = thought

        if first.font is not None and first.font.b and not status and not thoughts:
            sections.append({"name": title, "titles": []})
            continue

        entry = {"title": title, "status": status}
        if thoughts:
            entry["thoughts"] = thoughts
        sections[-1]["titles"].append(entry)

    return settings.get("key", label_key(ws.title).replace(" ", "-")), {
        "name": settings.get("name", ws.title),
        "people": [p["name"] for p in people],
        "sections": [s for s in sections if s["titles"]],
    }


def to_js(franchises, source):
    def dump(value):
        return json.dumps(value, ensure_ascii=False)

    lines = [
        "// Generated by tools/build_franchise_data.py from %s." % source,
        "// Don't edit this by hand: change the spreadsheet and run the script again.",
        "// Status keys: rewatched, watched, never, wont, notout. A person missing",
        "// from a title's status hasn't marked it yet.",
        "",
        "const FRANCHISES = {",
    ]
    for f_index, (key, franchise) in enumerate(franchises.items()):
        lines.append("  %s: {" % dump(key))
        lines.append('    "name": %s,' % dump(franchise["name"]))
        lines.append('    "people": %s,' % dump(franchise["people"]))
        lines.append('    "sections": [')
        for s_index, section in enumerate(franchise["sections"]):
            lines.append("      {")
            lines.append('        "name": %s,' % dump(section["name"]))
            lines.append('        "titles": [')
            for t_index, entry in enumerate(section["titles"]):
                lines.append("          {")
                lines.append('            "title": %s,' % dump(entry["title"]))
                thoughts = entry.get("thoughts")
                lines.append('            "status": %s%s' % (dump(entry["status"]), "," if thoughts else ""))
                if thoughts:
                    lines.append('            "thoughts": {')
                    items = list(thoughts.items())
                    for i, (name, thought) in enumerate(items):
                        lines.append("              %s: %s%s" % (dump(name), dump(thought), "," if i < len(items) - 1 else ""))
                    lines.append("            }")
                lines.append("          }" + ("," if t_index < len(section["titles"]) - 1 else ""))
            lines.append("        ]")
            lines.append("      }" + ("," if s_index < len(franchise["sections"]) - 1 else ""))
        lines.append("    ]")
        lines.append("  }" + ("," if f_index < len(franchises) - 1 else ""))
    lines.append("};")
    return "\n".join(lines) + "\n"


def main():
    site = Path(__file__).resolve().parent.parent
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else site / "Book.xlsx"
    wb = openpyxl.load_workbook(source)

    franchises = {}
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        key, franchise = read_sheet(ws)
        franchises[key] = franchise
        titles = sum(len(s["titles"]) for s in franchise["sections"])
        print("%s: %d titles, %d people (%s)" % (franchise["name"], titles, len(franchise["people"]), ", ".join(franchise["people"])))
        for section in franchise["sections"]:
            print("  %-32s %d" % (section["name"], len(section["titles"])))

    out = site / "franchise-data.js"
    out.write_text(to_js(franchises, source.name), encoding="utf-8", newline="\n")
    print("Wrote", out)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
